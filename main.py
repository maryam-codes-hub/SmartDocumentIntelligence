from pathlib import Path
from datetime import datetime
import fitz
from docx import Document as DocxDocument


# ================= DOCUMENT CLASS =================

class Document:

    def __init__(self,file):
        self.file = file
        self.extension = file.suffix.lower()

    def is_valid(self):
        return self.file.exists() and self.file.is_file()

    def extract_text(self):

        if self.extension == ".txt":
            return text_txt(self.file)

        elif self.extension == ".pdf":
            return text_pdf(self.file)

        elif self.extension == ".docx":
            return text_docx(self.file)

        else:
            return None

    def file_details(self):

        paragraph = "N/A"
        pages = "N/A"

        text = self.extract_text()

        if text is None:
            print("Unable to extract text.")
            return

        words = len(text.split())
        characters = len(text)

        # PDF pages
        if self.extension == ".pdf":

            doc = fitz.open(self.file)
            pages = len(doc)
            doc.close()

        # DOCX paragraphs
        elif self.extension == ".docx":

            doc = DocxDocument(self.file)
            paragraph = len(doc.paragraphs)

        details = {
            "name": self.file.name,
            "extension": self.extension,
            "size": format_file_size(self.file.stat().st_size),
            "modified": modified_time(self.file.stat().st_mtime),
            "Words": words,
            "characters": characters,
            "paragraphs": paragraph,
            "pages": pages
        }

        print("\n--------------------")
        print("Name:", details["name"])
        print("Type:", details["extension"])
        print("Size:", details["size"])
        print("Modified:", details["modified"])
        print("Words:", details["Words"])
        print("Characters:", details["characters"])
        print("Paragraphs:", details["paragraphs"])
        print("Pages:", details["pages"])


# ================= HELPER FUNCTIONS =================

def format_file_size(size):

    if size < 1024:
        return f"{size} B"

    elif size < 1024 ** 2:
        return f"{round(size / 1024, 2)} KB"

    elif size < 1024 ** 3:
        return f"{round(size / (1024 ** 2), 2)} MB"

    else:
        return f"{round(size / (1024 ** 3), 2)} GB"


def modified_time(timestamp):

    date_time = datetime.fromtimestamp(timestamp)

    return date_time.strftime("%d-%m-%Y %I:%M %p")


# ================= TEXT EXTRACTION FUNCTIONS =================

def text_txt(file):

    try:

        with open(file, "r", encoding="utf-8") as f:
            content = f.read()

        return content

    except UnicodeDecodeError:
        print("Unicode decode error.")

    except PermissionError:
        print("Permission denied.")


def text_pdf(file):

    content = []

    doc = fitz.open(file)

    for page in doc:
        content.append(page.get_text())

    doc.close()

    return "\n".join(content)


def text_docx(file):

    content = []

    doc = DocxDocument(file)

    for paragraph in doc.paragraphs:
        content.append(paragraph.text)

    return "\n".join(content)


# ================= DOCUMENT SCANNER =================

def scan_documents(folder_path):

    folder = Path(folder_path)

    if not folder.exists():
        print("Folder does not exist.")
        return

    if not folder.is_dir():
        print("Provided path is not a folder.")
        return

    supported_extensions = {".pdf", ".docx", ".txt"}

    documents = []

    try:

        for file in folder.rglob("*"):

            if file.is_file() and file.suffix.lower() in supported_extensions:

                documents.append({
                    "name": file.name,
                    "extension": file.suffix.lower(),
                    "size": format_file_size(file.stat().st_size),
                    "modified": modified_time(file.stat().st_mtime)
                })

    except PermissionError:
        print("Permission Denied.")

    print(f"\nDocuments found: {len(documents)}")

    for document in documents:

        print("\n--------------------")
        print("Name:", document["name"])
        print("Type:", document["extension"])
        print("Size:", document["size"])
        print("Modified:", document["modified"])


# ================= MAIN MENU =================

folder_path = input("Enter documents folder path: ")

while True:

    print("\n===== SmartDocumentIntelligence =====")
    print("1. View Documents")
    print("2. View File Content")
    print("3. View File Details")
    print("4. Exit")

    choice = input("Enter your choice: ")


    # ================= VIEW DOCUMENTS =================

    if choice == "1":

        scan_documents(folder_path)


    # ================= VIEW FILE CONTENT =================

    elif choice == "2":

        f_name = input("Enter File Name: ")

        file = Path(folder_path) / f_name

        document = Document(file)

        if not document.is_valid():
            print("File not found.")
            continue

        text = document.extract_text()

        if text is None:
            print("Unsupported file type.")
            continue

        print("\n========== FILE CONTENT ==========")
        print(text)


    # ================= VIEW FILE DETAILS =================

    elif choice == "3":

        print("\n==== File Details ====")

        f_name = input("Enter File Name: ")

        file = Path(folder_path) / f_name

        document = Document(file)

        if not document.is_valid():
            print("File not found.")
            continue

        document.file_details()


    # ================= EXIT =================

    elif choice == "4":

        print("Exiting SmartDocumentIntelligence...")
        break


    else:

        print("Invalid choice. Please try again.")